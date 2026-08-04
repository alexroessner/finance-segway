"""
Builds PUBLIC_FINANCE_template.xlsx — sovereign/municipal archetype (07).

Two independent lenses, since "public finance" spans two different
questions: is the issuer's overall debt load sustainable (debt sustainability
analysis, IMF/DSA-style), and can this specific revenue-backed bond cover its
own debt service (revenue bond coverage, muni-market-style). Both get their
own tab; neither depends on the other.
"""
import openpyxl
from template_helpers import *

wb = openpyxl.Workbook()
wb.remove(wb.active)

add_cover(wb, "[ISSUER] — Public Finance Model", [
    ("Issuer:", "[fill in — sovereign, state, muni, agency]"),
    ("Instrument:", "[general obligation / revenue bond / sovereign note]"),
    ("Last refreshed:", "[date]"),
    ("Next payment / issuance date:", "[date]"),
    ("Refresh cadence:", "Weekly"),
    ("Scenario (Base/Stress):", "Base"),
])

# ---------------- DEBT SUSTAINABILITY ----------------
ws = wb.create_sheet("Debt Sustainability")
set_col_widths(ws, [4, 40, 14, 46])
ws["B2"] = "Debt Sustainability Analysis (DSA)"; ws["B2"].font = TITLE
ws["B4"] = "Inputs"; ws["B4"].font = BOLD; ws["B4"].fill = GRAY_FILL
inputs = [
    ("Total public debt / GDP (or Debt/Revenue) ratio (%)", 0.60, PCT),
    ("Effective interest rate on debt, r (%)", 0.05, PCT2),
    ("Nominal GDP (or revenue base) growth rate, g (%)", 0.02, PCT2),
    ("Current primary balance (% of GDP; surplus positive)", 0.00, PCT2),
]
r = 5
for label, default, fmt in inputs:
    ws.cell(row=r, column=2, value=label).font = BLACK
    c = ws.cell(row=r, column=3, value=default); c.font = BLUE; c.fill = YELLOW_FILL
    c.number_format = fmt; c.border = BORDER
    r += 1

ws["B10"] = "Outputs"; ws["B10"].font = BOLD; ws["B10"].fill = GRAY_FILL
ws["B11"] = "Debt-stabilizing primary balance, pb* = (r-g)/(1+g) x debt ratio"
ws["C11"] = "=IFERROR((C6-C7)/(1+C7)*C5,\"-\")"
ws["C11"].font = BOLD; ws["C11"].number_format = PCT2; ws["C11"].border = BORDER
ws["D11"] = "Standard IMF/DSA formula. If r>g, debt ratio is explosive unless the primary balance is at least this large."
ws["D11"].font = ITALIC_GRAY

ws["B12"] = "Primary balance gap (current minus required)"
ws["C12"] = "=IFERROR(C8-C11,\"-\")"; ws["C12"].number_format = PCT2; ws["C12"].border = BORDER
ws["D12"] = "Negative = running a bigger deficit than needed to stabilize the debt ratio; ratio will rise."
ws["D12"].font = ITALIC_GRAY

ws["B13"] = "Debt trajectory"
ws["C13"] = '=IFERROR(IF(C8>=C11,"STABILIZING/FALLING","RISING"),"-")'; ws["C13"].font = BOLD
ws["D13"] = "Rule of thumb only — ignores stock-flow adjustments, FX-denominated debt, contingent liabilities."
ws["D13"].font = ITALIC_GRAY
ws.sheet_view.showGridLines = False

# ---------------- REVENUE BOND COVERAGE ----------------
ws = wb.create_sheet("Revenue Bond Coverage")
set_col_widths(ws, [4, 40, 14, 46])
ws["B2"] = "Revenue Bond Coverage"; ws["B2"].font = TITLE
ws["B4"] = "Inputs"; ws["B4"].font = BOLD; ws["B4"].fill = GRAY_FILL
inputs2 = [
    ("Gross pledged revenue ($)", 0, CUR),
    ("Operating & maintenance expense ($)", 0, CUR),
    ("Senior debt service — P&I ($)", 0, CUR),
    ("Subordinate debt service — P&I ($)", 0, CUR),
]
r = 5
for label, default, fmt in inputs2:
    ws.cell(row=r, column=2, value=label).font = BLACK
    c = ws.cell(row=r, column=3, value=default); c.font = BLUE; c.fill = YELLOW_FILL
    c.number_format = fmt; c.border = BORDER
    r += 1

ws["B10"] = "Outputs"; ws["B10"].font = BOLD; ws["B10"].fill = GRAY_FILL
ws["B11"] = "Net revenue available for debt service"
ws["C11"] = "=IFERROR(C5-C6,\"-\")"; ws["C11"].number_format = CUR; ws["C11"].border = BORDER
ws["B12"] = "Senior DSCR (Net revenue / Senior debt service)"
ws["C12"] = "=IFERROR(C11/C7,\"-\")"; ws["C12"].font = BOLD; ws["C12"].number_format = MULT
ws["C12"].border = BORDER
ws["B13"] = "All-in DSCR (Net revenue / Total debt service)"
ws["C13"] = "=IFERROR(C11/(C7+C8),\"-\")"; ws["C13"].font = BOLD; ws["C13"].number_format = MULT
ws["C13"].border = BORDER

ws["B15"] = "Additional Bonds Test covenant"; ws["B15"].font = BOLD; ws["B15"].fill = GRAY_FILL
ws["B16"] = "ABT minimum senior DSCR covenant"
c = ws.cell(row=16, column=3, value=1.25); c.font = BLUE; c.fill = YELLOW_FILL; c.number_format = MULT
c.border = BORDER
ws["B17"] = "Current-period headroom (actual minus covenant)"
ws["C17"] = "=IFERROR(C12-C16,\"-\")"; ws["C17"].number_format = MULT; ws["C17"].border = BORDER
ws["B18"] = "See the 'Additional Bonds Test' sheet for the forward-looking test on a proposed new issuance"
ws["B18"].font = ITALIC_GRAY
ws.sheet_view.showGridLines = False

# ---------------- ADDITIONAL BONDS TEST (forward-looking) ----------------
ws = wb.create_sheet("Additional Bonds Test")
set_col_widths(ws, [4, 44, 16, 16, 16, 18, 14, 12])
ws["B2"] = "Additional Bonds Test — Proposed New Issuance"; ws["B2"].font = TITLE
ws["B3"] = ("A static current-DSCR-vs-covenant check answers 'are we compliant today.' "
            "An ABT answers a different question: 'if we issue $X of new parity debt, do we STILL clear "
            "covenant' — the actual gate that decides whether a revenue-bond issuer can borrow more.")
ws["B3"].font = ITALIC_GRAY

ws["B5"] = "Proposed New Issuance"; ws["B5"].font = BOLD; ws["B5"].fill = GRAY_FILL
new_issuance_inputs = [
    ("New par amount ($)", 20_000_000, CUR),
    ("Coupon rate on new bonds (%)", 0.045, PCT2),
    ("Term of new bonds (years)", 20, NUM),
]
r = 6
for label, default, fmt in new_issuance_inputs:
    ws.cell(row=r, column=2, value=label).font = BLACK
    c = ws.cell(row=r, column=3, value=default); c.font = BLUE; c.fill = YELLOW_FILL
    c.number_format = fmt; c.border = BORDER
    r += 1
ws["B9"] = "New bond level annual debt service (P&I)"
ws["C9"] = "=IFERROR(-PMT(C7,C8,C6),\"-\")"; ws["C9"].font = BOLD; ws["C9"].number_format = CUR
ws["C9"].border = BORDER
ws["D9"] = "Assumes current-interest, level-debt-service structuring (the standard muni convention) — Excel's native PMT, not a capital-appreciation or non-level schedule."
ws["D9"].font = ITALIC_GRAY

ws["B11"] = "Existing Debt Service & Net Revenue"; ws["B11"].font = BOLD; ws["B11"].fill = GRAY_FILL
ws["B12"] = "Existing senior debt service ($)"
ws["C12"] = "='Revenue Bond Coverage'!C7"; ws["C12"].font = GREEN; ws["C12"].number_format = CUR
ws["C12"].border = BORDER
ws["B13"] = "Net revenue available for debt service ($)"
ws["C13"] = "='Revenue Bond Coverage'!C11"; ws["C13"].font = GREEN; ws["C13"].number_format = CUR
ws["C13"].border = BORDER
ws["B14"] = "ABT minimum senior DSCR covenant"
ws["C14"] = "='Revenue Bond Coverage'!C16"; ws["C14"].font = GREEN; ws["C14"].number_format = MULT
ws["C14"].border = BORDER

ws["B16"] = "Historical (Look-Back) Test"; ws["B16"].font = BOLD; ws["B16"].fill = GRAY_FILL
ws["B17"] = "Pro-forma total senior debt service ($) = existing + new"
ws["C17"] = "=C12+C9"; ws["C17"].number_format = CUR; ws["C17"].border = BORDER
ws["B18"] = "Pro-forma senior DSCR"
ws["C18"] = "=IFERROR(C13/C17,\"-\")"; ws["C18"].font = BOLD; ws["C18"].number_format = MULT
ws["C18"].border = BORDER
ws["B19"] = "Headroom (pro-forma DSCR minus covenant)"
ws["C19"] = "=IFERROR(C18-C14,\"-\")"; ws["C19"].number_format = MULT; ws["C19"].border = BORDER
ws["B20"] = "Historical test result"
ws["C20"] = '=IF(NOT(ISNUMBER(C18)),"-",IF(C18>=C14,"PASS","FAIL"))'; ws["C20"].font = BOLD

ws["B22"] = "Projected (Look-Forward) Test"; ws["B22"].font = BOLD; ws["B22"].fill = GRAY_FILL
ws["B23"] = "Active scenario (from Cover)"
ws["C23"] = "=Cover!C9"; ws["C23"].font = GREEN
scenario_inputs = [
    ("Revenue growth rate — Base (%)", 0.03, PCT2),
    ("Revenue growth rate — Stress (%)", -0.02, PCT2),
    ("O&M growth rate — Base (%)", 0.025, PCT2),
    ("O&M growth rate — Stress (%)", 0.05, PCT2),
]
r = 24
for label, default, fmt in scenario_inputs:
    ws.cell(row=r, column=2, value=label).font = BLACK
    c = ws.cell(row=r, column=3, value=default); c.font = BLUE; c.fill = YELLOW_FILL
    c.number_format = fmt; c.border = BORDER
    r += 1
ws["B28"] = "Active revenue growth rate"
ws["C28"] = '=IF(Cover!$C$9="Stress",C25,C24)'; ws["C28"].font = BOLD; ws["C28"].number_format = PCT2
ws["C28"].border = BORDER
ws["B29"] = "Active O&M growth rate"
ws["C29"] = '=IF(Cover!$C$9="Stress",C27,C26)'; ws["C29"].font = BOLD; ws["C29"].number_format = PCT2
ws["C29"].border = BORDER

headers = ["", "Year", "Gross Revenue", "O&M", "Net Revenue", "Pro-forma Senior DS", "Pro-forma DSCR", "Pass/Fail"]
for i, h in enumerate(headers, start=1):
    ws.cell(row=31, column=i, value=h)
style_header_row(ws, 31, 7, start_col=2)
for yr in range(1, 6):
    row = 31 + yr
    ws.cell(row=row, column=2, value=f"Year {yr}").font = BOLD
    if yr == 1:
        gross_prev = "'Revenue Bond Coverage'!$C$5"
        om_prev = "'Revenue Bond Coverage'!$C$6"
    else:
        gross_prev = f"C{row - 1}"
        om_prev = f"D{row - 1}"
    ws.cell(row=row, column=3, value=f"={gross_prev}*(1+$C$28)").number_format = CUR
    ws.cell(row=row, column=4, value=f"={om_prev}*(1+$C$29)").number_format = CUR
    ws.cell(row=row, column=5, value=f"=C{row}-D{row}").number_format = CUR
    ws.cell(row=row, column=6, value="=$C$17").number_format = CUR
    ws.cell(row=row, column=7, value=f"=IFERROR(E{row}/F{row},\"-\")").number_format = MULT
    ws.cell(row=row, column=7).font = BOLD
    ws.cell(row=row, column=8, value=f'=IF(G{row}>=$C$14,"PASS","FAIL")')
    for col in range(3, 9):
        ws.cell(row=row, column=col).border = BORDER

ws["B38"] = "Minimum projected DSCR (5-yr window)"
ws["C38"] = "=MIN(G32:G36)"; ws["C38"].font = BOLD; ws["C38"].number_format = MULT
ws["C38"].border = BORDER
ws["B39"] = "Year of minimum"
ws["C39"] = '=INDEX(B32:B36,MATCH(C38,G32:G36,0))'; ws["C39"].border = BORDER
ws["B40"] = "Projected test result"
ws["C40"] = '=IF(C38>=C14,"PASS — all 5 years clear covenant","FAIL — covenant breached in "&C39)'
ws["C40"].font = BOLD

ws["B42"] = ("Indentures typically require passing the historical OR the projected test (not always both) before "
             "additional parity debt may be issued — confirm which applies for a specific issuer's indenture; "
             "see model_card.md.")
ws["B42"].font = ITALIC_GRAY
ws.sheet_view.showGridLines = False

add_sources_checks(
    wb,
    sources=[
        ("Debt-stabilizing primary balance, pb* = (r-g)/(1+g) x debt ratio", "IMF Debt Sustainability Analysis (DSA) framework", "Standard practice", "Public-domain IMF methodology, not issuer-specific"),
        ("Additional Bonds Test: historical + projected look-forward tests", "Typical municipal revenue-bond indenture additional-parity-debt covenant", "Standard practice", "Actual indenture language varies by issuer; implements the common historical-or-projected pattern"),
        ("New-issuance level debt service (PMT amortization)", "Standard level-debt-service muni bond structuring convention", "Standard practice", "Current-interest, level P&I only — no capital-appreciation or non-level schedules"),
        ("ABT minimum senior DSCR covenant (1.25x default)", "[fill in — issuer's actual bond indenture]", "[fill in]", "Placeholder default; replace with the specific indenture's real covenant level before use"),
    ],
    checks=[
        ("Net revenue ties: Gross revenue - O&M = Net revenue", "='Revenue Bond Coverage'!C11-('Revenue Bond Coverage'!C5-'Revenue Bond Coverage'!C6)", "0 (exact)"),
        ("Pro-forma Yr1 debt service ties: existing + new = pro forma", "='Additional Bonds Test'!C17-('Additional Bonds Test'!C12+'Additional Bonds Test'!C9)", "0 (exact)"),
        ("New-bond PMT reproduces closed-form level-annuity payment", "='Additional Bonds Test'!C9-('Additional Bonds Test'!C6*'Additional Bonds Test'!C7/(1-(1+'Additional Bonds Test'!C7)^-'Additional Bonds Test'!C8))", "~0 (native PMT matches closed-form annuity formula)"),
        ("DSA trajectory label matches sign of primary-balance gap", '=IF(\'Debt Sustainability\'!C12>=0,"STABILIZING/FALLING","RISING")=\'Debt Sustainability\'!C13', "TRUE"),
    ],
)

add_refresh_log(wb)
out_path = "PUBLIC_FINANCE_template.xlsx"
wb.save(out_path)
print("saved", out_path)
