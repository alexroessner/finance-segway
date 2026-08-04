"""
postgres_etl.py — extract verified LBO/PE deal outputs into Postgres.

This is a query/reporting layer, not a second calculation engine: every
number loaded here comes from a workbook that has ALREADY been recalculated
for real by tools/recalc.py (headless LibreOffice), read back with openpyxl
data_only=True. Nothing in this script re-implements an IRR solve, a cash
sweep cascade, or a returns waterfall in Python or SQL — it reuses the same
toggle-and-recalc mechanism already independently verified by
check_lbo_scenario_switch in tools/verify_reference_calcs.py (set
Cover!C11 to "Base" or "Downside", recalc, read the Active column).

Usage:
    python3 tools/postgres_etl.py                 # load all registered deals
    python3 tools/postgres_etl.py --dsn "dbname=finance_segway"
    python3 tools/postgres_etl.py --dry-run        # print extracted rows, skip DB writes

Requires: psycopg2 (pip install psycopg2-binary), a reachable Postgres
database with db/schema.sql already applied.
"""
import argparse
import os
import shutil
import sys
import tempfile

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO_ROOT, "tools"))

import openpyxl  # noqa: E402
from recalc import recalc  # noqa: E402

# Deal registry: workbook path (repo-relative) -> sector tag. Sector isn't
# a field in the workbook itself (Cover only has sponsor/deal type/dates),
# so it's tracked here, matching the sector-playbook pilot in
# 03_Private_Equity/playbooks/. Add a new deal by adding one line here.
DEALS = {
    "PROJECT_ATLAS": {
        "path": "03_Private_Equity/deals/PROJECT_ATLAS.xlsx",
        "sector": "Diversified / Illustrative",
    },
    "HEALTHCARE_SERVICES_EXAMPLE": {
        "path": "03_Private_Equity/deals/HEALTHCARE_SERVICES_EXAMPLE.xlsx",
        "sector": "Healthcare Services",
    },
    "INDUSTRIALS_EXAMPLE": {
        "path": "03_Private_Equity/deals/INDUSTRIALS_EXAMPLE.xlsx",
        "sector": "Industrials",
    },
    "SAAS_EXAMPLE": {
        "path": "03_Private_Equity/deals/SAAS_EXAMPLE.xlsx",
        "sector": "SaaS",
    },
}

SCENARIOS = ["Base", "Downside"]

# Debt Schedule tab: (label, row) for the per-year series read from the
# Active column (columns C..H = Yr0..Yr5).
DEBT_SCHEDULE_ROWS = {
    "ebitda": 5,
    "fcf": 6,
    "revolver_ending_balance": 13,
    "tla_ending_balance": 20,
    "tlb_ending_balance": 27,
    "total_debt_ending": 30,
    "total_interest_expense": 31,
    "leverage_multiple": 32,
}


def recalc_with_scenario(src_path, scenario):
    """Copy src_path to a temp file, set Cover!C11 to the requested
    scenario, recalc for real, and return the recalculated (data_only)
    workbook. Mirrors with_recalc() in tools/verify_reference_calcs.py."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    shutil.copy(src_path, tmp.name)
    try:
        wb = openpyxl.load_workbook(tmp.name)
        wb["Cover"]["C11"] = scenario
        wb.save(tmp.name)
        result = recalc(tmp.name, timeout=45)
        if "error" in result:
            raise RuntimeError(f"recalc failed for {src_path} ({scenario}): {result['error']}")
        if result.get("status") != "success":
            raise RuntimeError(f"recalc found formula errors for {src_path} ({scenario}): {result}")
        return openpyxl.load_workbook(tmp.name, data_only=True)
    finally:
        os.unlink(tmp.name)


def extract_cover(wb):
    cov = wb["Cover"]
    return {
        "deal_name": str(cov["B2"].value or "").replace("[TARGET] — ", "").strip() or None,
        "sponsor": cov["C4"].value,
        "deal_type": cov["C5"].value,
        "entry_date": cov["C8"].value,
        "hold_period_years": cov["C9"].value,
    }


def extract_sources_uses(wb):
    su = wb["Sources & Uses"]
    rows = []
    r = 5
    while su.cell(row=r, column=2).value and su.cell(row=r, column=2).value not in ("Total sources",):
        label = su.cell(row=r, column=2).value
        amount = su.cell(row=r, column=3).value
        if isinstance(amount, (int, float)):
            rows.append(("Sources", label, amount))
        r += 1
    r = 5
    while su.cell(row=r, column=5).value and su.cell(row=r, column=5).value not in ("Total uses",):
        label = su.cell(row=r, column=5).value
        amount = su.cell(row=r, column=6).value
        if isinstance(amount, (int, float)):
            rows.append(("Uses", label, amount))
        r += 1
    return rows


def extract_returns(wb):
    ret = wb["Returns"]
    ds = wb["Debt Schedule"]
    return {
        "entry_ebitda": ret["C5"].value,
        "entry_multiple": ret["C6"].value,
        "entry_ev": ret["C7"].value,
        "sponsor_equity_entry": ret["C8"].value,
        "exit_ebitda": ret["C11"].value,
        "exit_multiple": ret["C12"].value,
        "exit_ev": ret["C13"].value,
        "net_debt_at_exit": ret["C14"].value,
        "exit_equity_value": ret["C15"].value,
        "moic_blended": ret["C18"].value,
        "irr_blended": ret["C20"].value,
        "hold_period_years": ret["C19"].value,
        "management_promote": ret["C29"].value,
        "sponsor_net_proceeds": ret["C31"].value,
        "sponsor_moic": ret["C32"].value,
        "sponsor_irr": ret["C33"].value,
        "yr5_leverage": ds.cell(row=32, column=8).value,  # column H = Yr5
    }


def extract_debt_schedule(wb):
    ds = wb["Debt Schedule"]
    rows = []
    for year in range(6):  # Yr0..Yr5
        col = 3 + year  # C=Yr0 .. H=Yr5
        row = {"year": year}
        for field, r in DEBT_SCHEDULE_ROWS.items():
            row[field] = ds.cell(row=r, column=col).value
        rows.append(row)
    return rows


def extract_deal(deal_id, path):
    """Recalc under both scenarios and return the full extracted payload."""
    abs_path = os.path.join(REPO_ROOT, path)
    payload = {"deal_id": deal_id, "path": path, "returns": {}, "debt_schedule": {}}
    for i, scenario in enumerate(SCENARIOS):
        wb = recalc_with_scenario(abs_path, scenario)
        if i == 0:
            payload["cover"] = extract_cover(wb)
            payload["sources_uses"] = extract_sources_uses(wb)
        payload["returns"][scenario] = extract_returns(wb)
        payload["debt_schedule"][scenario] = extract_debt_schedule(wb)
    return payload


def load_deal(cur, deal_id, sector, payload):
    cov = payload["cover"]
    cur.execute(
        """
        INSERT INTO deals (deal_id, deal_name, sponsor, deal_type, sector,
                            entry_date, hold_period_years, workbook_path, last_synced_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, now())
        ON CONFLICT (deal_id) DO UPDATE SET
            deal_name = EXCLUDED.deal_name, sponsor = EXCLUDED.sponsor,
            deal_type = EXCLUDED.deal_type, sector = EXCLUDED.sector,
            entry_date = EXCLUDED.entry_date, hold_period_years = EXCLUDED.hold_period_years,
            workbook_path = EXCLUDED.workbook_path, last_synced_at = now()
        """,
        (deal_id, cov["deal_name"], cov["sponsor"], cov["deal_type"], sector,
         cov["entry_date"], cov["hold_period_years"], payload["path"]),
    )

    cur.execute("DELETE FROM deal_sources_uses WHERE deal_id = %s", (deal_id,))
    for side, label, amount in payload["sources_uses"]:
        cur.execute(
            "INSERT INTO deal_sources_uses (deal_id, side, line_item, amount) VALUES (%s, %s, %s, %s)",
            (deal_id, side, label, amount),
        )

    for scenario in SCENARIOS:
        r = payload["returns"][scenario]
        cur.execute(
            """
            INSERT INTO deal_returns (deal_id, scenario, entry_ebitda, entry_multiple, entry_ev,
                sponsor_equity_entry, exit_ebitda, exit_multiple, exit_ev, net_debt_at_exit,
                exit_equity_value, moic_blended, irr_blended, hold_period_years,
                management_promote, sponsor_net_proceeds, sponsor_moic, sponsor_irr, yr5_leverage)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (deal_id, scenario) DO UPDATE SET
                entry_ebitda = EXCLUDED.entry_ebitda, entry_multiple = EXCLUDED.entry_multiple,
                entry_ev = EXCLUDED.entry_ev, sponsor_equity_entry = EXCLUDED.sponsor_equity_entry,
                exit_ebitda = EXCLUDED.exit_ebitda, exit_multiple = EXCLUDED.exit_multiple,
                exit_ev = EXCLUDED.exit_ev, net_debt_at_exit = EXCLUDED.net_debt_at_exit,
                exit_equity_value = EXCLUDED.exit_equity_value, moic_blended = EXCLUDED.moic_blended,
                irr_blended = EXCLUDED.irr_blended, hold_period_years = EXCLUDED.hold_period_years,
                management_promote = EXCLUDED.management_promote,
                sponsor_net_proceeds = EXCLUDED.sponsor_net_proceeds,
                sponsor_moic = EXCLUDED.sponsor_moic, sponsor_irr = EXCLUDED.sponsor_irr,
                yr5_leverage = EXCLUDED.yr5_leverage
            """,
            (deal_id, scenario, r["entry_ebitda"], r["entry_multiple"], r["entry_ev"],
             r["sponsor_equity_entry"], r["exit_ebitda"], r["exit_multiple"], r["exit_ev"],
             r["net_debt_at_exit"], r["exit_equity_value"], r["moic_blended"], r["irr_blended"],
             r["hold_period_years"], r["management_promote"], r["sponsor_net_proceeds"],
             r["sponsor_moic"], r["sponsor_irr"], r["yr5_leverage"]),
        )

        cur.execute("DELETE FROM deal_debt_schedule WHERE deal_id = %s AND scenario = %s",
                    (deal_id, scenario))
        for row in payload["debt_schedule"][scenario]:
            cur.execute(
                """
                INSERT INTO deal_debt_schedule (deal_id, scenario, year, ebitda, fcf,
                    revolver_ending_balance, tla_ending_balance, tlb_ending_balance,
                    total_debt_ending, total_interest_expense, leverage_multiple)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (deal_id, scenario, row["year"], row["ebitda"], row["fcf"],
                 row["revolver_ending_balance"], row["tla_ending_balance"],
                 row["tlb_ending_balance"], row["total_debt_ending"],
                 row["total_interest_expense"], row["leverage_multiple"]),
            )


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dsn", default=os.environ.get("FINANCE_SEGWAY_PG_DSN", "dbname=finance_segway"))
    ap.add_argument("--dry-run", action="store_true", help="extract and print, skip DB writes")
    args = ap.parse_args()

    conn = None
    if not args.dry_run:
        import psycopg2
        conn = psycopg2.connect(args.dsn)

    try:
        for deal_id, meta in DEALS.items():
            print(f"Extracting {deal_id} ({meta['path']}) ...", file=sys.stderr)
            payload = extract_deal(deal_id, meta["path"])
            if args.dry_run:
                for scenario in SCENARIOS:
                    r = payload["returns"][scenario]
                    print(f"  {scenario}: sponsor IRR={r['sponsor_irr']:.4f} "
                          f"sponsor MOIC={r['sponsor_moic']:.4f} Yr5 leverage={r['yr5_leverage']:.3f}x")
                continue
            with conn.cursor() as cur:
                load_deal(cur, deal_id, meta["sector"], payload)
            conn.commit()
            print(f"  loaded {deal_id} into Postgres", file=sys.stderr)
    finally:
        if conn is not None:
            conn.close()

    print("Done." if not args.dry_run else "Dry run complete (no DB writes).", file=sys.stderr)


if __name__ == "__main__":
    main()
